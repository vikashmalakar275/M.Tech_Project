from __future__ import annotations

import json
import shutil
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from orbitwatch.experiment import resolve_run
from orbitwatch.service import TelemetryService


def create_submission(root: Path, run: str | None = None) -> Path:
    run_dir = resolve_run(root, run)
    manifest = json.loads((run_dir / "manifest.json").read_text())
    summary = json.loads((run_dir / "summary.json").read_text())
    metrics = pd.read_csv(run_dir / "metrics.csv")
    output = root / "submission"
    output.mkdir(exist_ok=True)
    for source, name in (
        (run_dir / "metrics.csv", "detection_metrics.csv"),
        (run_dir / "summary.json", "detection_summary.json"),
        (run_dir / "manifest.json", "experiment_manifest.json"),
    ):
        shutil.copyfile(source, output / name)
    explanation_path = run_dir / "explanations" / "summary.json"
    explanations = json.loads(explanation_path.read_text()) if explanation_path.exists() else None
    if explanations:
        for name in ("summary.json", "cases.csv", "records.json"):
            shutil.copyfile(run_dir / "explanations" / name, output / f"explanation_{name}")
    all_results = [item for item in summary if item["spacecraft"] == "ALL"]
    channel_count = len(manifest["channels"])
    samples = int(metrics.loc[metrics["detector"] == "gru", "evaluated_samples"].sum())
    statements = [
        f"Evaluated {channel_count} independent NASA channel files and {samples:,} post-warmup test observations.",
        "Two detectors: causal rolling-median forecasting and a small GRU; no published architecture is mislabelled.",
        "The final 20% of each supplied training sequence calibrates residual thresholds. Test labels are evaluation-only.",
        "MCP exposes read-only research tools. A restricted claim validator rejects unsupported values and citations.",
        "All results are generated from saved experiment artifacts; no target accuracy is assumed.",
    ]
    report_title = "OrbitWatch: Evidence-Grounded Spacecraft Telemetry Investigation"
    document = SimpleDocTemplate(
        str(output / "Technical_Report.pdf"),
        pagesize=A4,
        rightMargin=1.8 * cm,
        leftMargin=1.8 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.8 * cm,
        title=report_title,
        author="OrbitWatch project",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="SmallBody", parent=styles["BodyText"], fontSize=8, leading=11))
    story = []

    def paragraph(text: str, style: str = "BodyText") -> None:
        story.append(Paragraph(escape(text), styles[style]))
        story.append(Spacer(1, 0.18 * cm))

    def heading(text: str) -> None:
        paragraph(text, "Heading2")

    def table(rows: list[list[str]], widths: list[float] | None = None) -> None:
        converted = [
            [Paragraph(escape(str(value)), styles["SmallBody"]) for value in row] for row in rows
        ]
        item = Table(converted, colWidths=widths, repeatRows=1, hAlign="LEFT")
        item.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dceef5")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#c5d1db")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.extend([item, Spacer(1, 0.3 * cm)])

    paragraph(report_title, "Title")
    paragraph(
        "Final-semester project technical report - generated from measured results", "Heading2"
    )
    paragraph(f"Experiment: {manifest['run_id']} | Completed: {manifest['completed_at']}")
    paragraph(
        "Submission draft for supervisor review. This is not an approved thesis, a publication claim, or a certification of mission readiness."
    )
    heading("Abstract")
    paragraph(
        f"This project implements a locally runnable spacecraft-telemetry investigation workbench. "
        f"It evaluates two lightweight detectors on {channel_count} anonymized NASA SMAP/MSL streams "
        f"({samples:,} evaluated samples), with chronological calibration and no point adjustment. "
        "Detected intervals become structured evidence exposed through Model Context Protocol tools. "
        "A deterministic report and a local-language-model report share a numerical validation layer. "
        "The work separates detection performance from explanation correctness, and preserves negative "
        "results and dataset limitations rather than claiming unsupported physical diagnosis."
    )
    heading("1. Research question and contribution")
    paragraph(
        "Can evidence-constrained investigation reject unsupported numerical and causal claims while retaining useful factual coverage?"
    )
    for statement in statements:
        paragraph(statement)
    paragraph(
        "The contribution is a reproducible systems integration and controlled explanation comparison, not a new GRU architecture or a state-of-the-art detector. Related LLM/time-series work must be discussed before making any novelty claim."
    )
    heading("2. System architecture")
    paragraph(
        "Recorded NASA files -> channel-isolated chronological split -> forecasting -> calibrated residual threshold -> event evidence -> MCP server/client -> claim proposal -> validation -> dashboard and cited report."
    )
    paragraph(
        "The Streamlit application replays saved causal scores. Each forecast uses observations strictly before its target. The replay cursor limits visible samples, alerts, and evidence. This is not a live mission feed."
    )
    heading("3. Data and provenance")
    paragraph(manifest["dataset"].get("caveat", "See dataset manifest for source limitations."))
    if "annotation_policy" in manifest["dataset"]:
        paragraph(manifest["dataset"]["annotation_policy"])
    paragraph(
        f"Dataset archive SHA-256: {manifest['dataset'].get('archive_sha256', 'not recorded')}",
        "SmallBody",
    )
    paragraph(f"Computational source SHA-256: {manifest['source_sha256']}", "SmallBody")
    paragraph(
        "Ground-truth intervals are inclusive sample indices. Channel files are never concatenated into one timeline. The first feature is the target telemetry; remaining features are encoded commands, not a verified synchronized spacecraft sensor graph."
    )
    heading("4. Detection methodology")
    config = manifest["config"]
    paragraph(
        f"Context window: {config['window']} samples; GRU hidden width: {config['hidden']}; "
        f"epochs: {config['epochs']}; seed: {config['seed']}; maximum fitting windows/channel: "
        f"{config['max_train_windows']}. Training uses Smooth L1 loss and AdamW. "
        "The standardization mean and scale are fitted only on the first training segment. "
        "The last calibration_fraction of the supplied training file is not used for weight fitting."
    )
    paragraph(
        f"Residual score s(t) = |normalized observation(t) - normalized prediction(t)|. "
        f"The fixed threshold is calibration quantile {config['threshold_quantile']}. "
        "The statistical baseline predicts the median of the previous context window. "
        "The test stream has its own warmup and is not joined to the training file."
    )
    heading("5. Detection results")
    table(
        [["Detector", "Macro point F1", "Macro AP", "Event recall", "FP / 1000 normal"]]
        + [
            [
                r["detector"],
                f"{r['macro_point_f1']:.4f}",
                f"{r['macro_average_precision']:.4f}",
                f"{r['event_recall']:.4f}",
                f"{r['false_alarms_per_1000_normal_samples']:.2f}",
            ]
            for r in all_results
        ],
        [3.3 * cm, 3 * cm, 2.5 * cm, 3 * cm, 4 * cm],
    )
    paragraph(
        "Macro point F1 and average precision give each channel equal weight. Event recall pools matched and true intervals. Event matching is chronological, one-to-one temporal overlap; fragmented alarms incur extra predicted events. False-positive rates accompany event metrics to discourage overly long alarms. Detection delay is in samples, not seconds. The CSV retains every channel and missed event contribution."
    )
    winner = max(all_results, key=lambda item: item["macro_point_f1"])
    paragraph(
        f"In this run, {winner['detector']} has the higher macro point F1 ({winner['macro_point_f1']:.4f}). This is an observation under the declared configuration, not evidence of universal superiority. No hyperparameter selection using test labels is claimed."
    )
    heading("6. Evidence-constrained explanation")
    paragraph(
        "Permitted claims cover the channel, event start/end, peak sample, peak ratio, observed and predicted values, and the explicit unknown-cause boundary. Each claim must cite the current evidence ID. Indices and strings match exactly; measured floating values match within 1e-6. Unsupported and duplicate claims are rejected and displayed separately."
    )
    paragraph(
        "Final prose is rendered from validated fact fields, not copied from unrestricted model text. This intentionally limits expressiveness. Correctness within this schema does not establish semantic truthfulness of arbitrary LLM explanations."
    )
    if explanations:
        paragraph(f"The local pilot used {explanations['model']}. {explanations['selection']}")
        table(
            [["Mode", "Cases", "Unsupported emitted", "Coverage", "Mean seconds"]]
            + [
                [
                    r["mode"],
                    str(r["cases"]),
                    f"{r['unsupported_emitted_rate']:.3f}"
                    if r["unsupported_emitted_rate"] is not None
                    else "n/a",
                    f"{r['mean_coverage']:.3f}",
                    f"{r['mean_latency_seconds']:.2f}",
                ]
                for r in explanations["summary"]
            ],
            [5.2 * cm, 1.3 * cm, 3.2 * cm, 2.1 * cm, 4 * cm],
        )
        for limitation in explanations["limitations"]:
            paragraph(limitation, "SmallBody")
    else:
        paragraph(
            "No local-model explanation experiment exists for this run. Explanation metrics are therefore not reported."
        )
    heading("7. Reproducibility and implementation")
    paragraph(
        "README.md provides installation, training, replay, MCP, and report commands. Pinned direct dependencies and requirements-lock.txt record the software environment. The repository includes tests for temporal causality, calibration isolation, model reload, metric definitions, evidence checks, local history, actual MCP subprocess exchange, and application rendering."
    )
    paragraph(
        "MCP uses the pinned 1.x Python SDK over local stdio. The application calls tools through a real client session, rather than simulating a tool trace. All dataset artifacts, local models, and SQLite records stay outside version control."
    )
    heading("8. Limitations and appropriate claims")
    for limitation in manifest["limitations"]:
        paragraph(limitation)
    paragraph(
        "The explanation pilot is small and uses machine-checkable facts, not expert-labelled physical diagnoses or an operator user study. A single seed does not support confidence intervals. Anonymous benchmark preprocessing limits deployment conclusions. No hardware-energy measurement or certified onboard feasibility is claimed."
    )
    heading("9. Conclusion and future work")
    paragraph(
        "OrbitWatch delivers an end-to-end, local, traceable research prototype and a reproducible comparison. Its primary value is the separation of numerical evidence from unsupported interpretation. Extensions should prioritize repeated-seed evaluation, independent expert annotation, operationally meaningful datasets, and calibration robustness rather than adding unvalidated model names."
    )
    heading("References")
    for reference in (
        "Hundman et al. Detecting Spacecraft Anomalies Using LSTMs and Nonparametric Dynamic Thresholding. KDD 2018. https://arxiv.org/abs/1802.04431",
        "Kim et al. Towards a Rigorous Evaluation of Time-series Anomaly Detection. https://arxiv.org/abs/2109.05257",
        "Alnegheimish et al. Can Large Language Models be Anomaly Detectors for Time Series? DSAA 2024. https://arxiv.org/abs/2405.14755",
        "NASA benchmark implementation and data documentation: https://github.com/khundman/telemanom",
        "Model Context Protocol Python SDK: https://github.com/modelcontextprotocol/python-sdk",
    ):
        paragraph(reference, "SmallBody")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(1.8 * cm, 1.1 * cm, "OrbitWatch | Generated measured-results report")
        canvas.drawRightString(A4[0] - 1.8 * cm, 1.1 * cm, str(doc.page))
        canvas.restoreState()

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    def slide(title: str, lines: list[str]) -> None:
        item = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = item.background.fill
        background.solid()
        background.fore_color.rgb = RGBColor(8, 17, 31)
        title_box = item.shapes.add_textbox(
            Inches(0.65), Inches(0.55), Inches(12), Inches(1)
        ).text_frame
        title_box.text = title
        title_box.paragraphs[0].font.size = Pt(32)
        title_box.paragraphs[0].font.bold = True
        title_box.paragraphs[0].font.color.rgb = RGBColor(69, 219, 179)
        body = item.shapes.add_textbox(
            Inches(0.75), Inches(1.75), Inches(11.8), Inches(5)
        ).text_frame
        body.word_wrap = True
        for index, line in enumerate(lines):
            p = body.paragraphs[0] if index == 0 else body.add_paragraph()
            p.text = line
            p.font.size = Pt(21)
            p.font.color.rgb = RGBColor(220, 231, 243)
            p.space_after = Pt(20)
        notes = item.notes_slide.notes_text_frame
        notes.text = "Generated from actual experiment artifacts. Explain limitations and avoid claiming physical root-cause diagnosis."

    slide(
        "OrbitWatch",
        [
            "Evidence-Grounded Spacecraft Telemetry Investigation",
            "Final-semester research prototype",
            f"Measured experiment: {manifest['run_id']}",
            "Draft for supervisor review; not a publication or flight-readiness claim.",
        ],
    )
    slide(
        "Continuity with earlier work",
        [
            "Earlier project: spacecraft telemetry anomaly detection.",
            "Summer study: Model Context Protocol integration.",
            "Final deliverable: working detectors + traceable MCP investigation + measured results.",
        ],
    )
    slide(
        "Focused research question",
        [
            "Can evidence validation reject unsupported claims while retaining factual coverage?",
            "Separate detection accuracy from explanation correctness.",
            "A systems/experimental contribution, not a new neural architecture.",
        ],
    )
    slide(
        "System architecture",
        [
            "Recorded NASA data -> causal forecasters -> fixed thresholds.",
            "Alert intervals -> structured evidence -> MCP tools.",
            "Template or local LLM -> claim checks -> cited report.",
            "Streamlit console, replay controls, local SQLite history.",
        ],
    )
    slide(
        "Data and evaluation discipline",
        [
            f"{channel_count} channel files; {samples:,} evaluated observations.",
            "Chronological fitting/calibration; never concatenate channels.",
            "No point adjustment; record false alarms, delay, and missed events.",
            "Anonymous data and upstream test-extrema scaling are explicit limitations.",
        ],
    )
    slide(
        "Lightweight detectors",
        [
            f"Rolling median versus GRU (hidden={config['hidden']}, window={config['window']}).",
            f"{config['epochs']} epochs; seed={config['seed']}; CPU execution.",
            f"Threshold: calibration quantile {config['threshold_quantile']}.",
            "Simple baselines may win; report measured outcomes honestly.",
        ],
    )
    slide(
        "Measured detection results",
        [
            f"{r['detector']}: macro F1={r['macro_point_f1']:.4f}; AP={r['macro_average_precision']:.4f}; event recall={r['event_recall']:.4f}."
            for r in all_results
        ]
        + [
            "Full per-channel measurements are included in detection_metrics.csv.",
            "These scores are not comparable to point-adjusted F1 tables.",
        ],
    )
    slide(
        "Evidence instead of invented diagnosis",
        [
            "Each fact cites an evidence ID and a recorded numerical value.",
            "Unknown fields, incorrect values, and invalid citations are rejected.",
            "Physical root cause remains unknown.",
            "Trusted rendering limits expressiveness but makes the boundary inspectable.",
        ],
    )
    if explanations:
        slide(
            "Local explanation pilot",
            [f"Model: {explanations['model']}; paired complete/redacted evidence."]
            + [
                f"{r['mode']}: unsupported={r['unsupported_emitted_rate']:.3f}; coverage={r['mean_coverage']:.3f}."
                for r in explanations["summary"]
            ],
        )
    slide(
        "Live demonstration",
        [
            "Replay one channel and reveal a predicted alert.",
            "Inspect observations, forecasts, and threshold ratios.",
            "Investigate through a real MCP session and inspect its trace.",
            "Export cited evidence; demonstrate refusal to invent a hardware cause.",
        ],
    )
    slide(
        "Limitations and next experiments",
        [
            "Single-seed detector study and a small claim-level explanation pilot.",
            "No expert operator study or physical cause labels.",
            "No claim of state-of-the-art or onboard certification.",
            "Next: repeated seeds, expert evaluation, and external operational datasets.",
        ],
    )
    slide(
        "Deliverables",
        [
            "Executable local dashboard and MCP server.",
            "Pinned environment, automated tests, and reproducible commands.",
            "Trained local artifacts plus transparent measured result tables.",
            "Technical report, presentation, and GitHub source.",
        ],
    )
    presentation.save(output / "Final_Presentation.pptx")
    service = TelemetryService(root, run_dir.name)
    demo_events = []
    for channel in manifest["channels"]:
        events = service.events(channel)
        if events:
            best = max(events, key=lambda item: item["peak_ratio"])
            demo_events.append(
                {
                    "channel": channel,
                    "spacecraft": service.channel_metadata(channel)["spacecraft"],
                    **best,
                }
            )
    (output / "demo_events.json").write_text(json.dumps(demo_events, indent=2))
    return output
