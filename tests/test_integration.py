from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation
from streamlit.testing.v1 import AppTest

from orbitwatch.evidence import Claim, ClaimDraft, template_report
from orbitwatch.explanation_eval import evaluate_explanations
from orbitwatch.mcp_client import investigate_via_mcp
from orbitwatch.service import TelemetryService
from orbitwatch.submission import create_submission


@pytest.mark.asyncio
async def test_actual_mcp_subprocess_exchange(completed_root):
    service = TelemetryService(completed_root)
    event = max(service.events("A-1"), key=lambda item: item["peak_ratio"])
    evidence, trace = await investigate_via_mcp(
        completed_root, "A-1", event["start"], event["end"], "gru", 159, "fixture"
    )
    assert evidence.facts["channel"] == "A-1"
    assert "get_event_evidence" in trace[0]["tools"]
    assert trace[1]["evidence_id"] == evidence.evidence_id


def test_streamlit_renders_and_changes_detector(completed_root, monkeypatch):
    monkeypatch.setenv("ORBITWATCH_HOME", str(completed_root))
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py"), default_timeout=30).run()
    assert not app.exception
    assert len(app.metric) == 4
    app.selectbox(key="detector").select("rolling_median").run()
    assert not app.exception


def test_streamlit_mcp_investigation_and_replay(completed_root, monkeypatch):
    monkeypatch.setenv("ORBITWATCH_HOME", str(completed_root))
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py"), default_timeout=30).run()
    app.button(key="investigate").click().run(timeout=60)
    assert not app.exception
    assert not app.error
    assert any("OrbitWatch investigation" in element.value for element in app.markdown)
    app.button(key="restart").click().run()
    assert not app.exception
    assert app.metric[0].value == "9"
    app.button(key="finish").click().run()
    assert app.metric[0].value == "160"


def test_streamlit_empty_install_is_actionable(tmp_path, monkeypatch):
    monkeypatch.setenv("ORBITWATCH_HOME", str(tmp_path))
    app = AppTest.from_file(str(Path(__file__).parents[1] / "app.py"), default_timeout=30).run()
    assert not app.exception
    assert "Download the benchmark" in app.info[0].value


def test_results_have_real_provenance(completed_root):
    manifest = json.loads((completed_root / "runs/fixture/manifest.json").read_text())
    assert manifest["status"] == "complete"
    assert len(manifest["source_sha256"]) == 64
    assert manifest["config"]["seed"] == 17
    assert manifest["duration_seconds"] > 0


@pytest.mark.parametrize("invalid_draft", [False, True])
def test_explanation_timing_and_nullable_submission(completed_root, monkeypatch, invalid_draft):
    def draft(evidence, model, *, grounded):
        claims = (
            [
                Claim(
                    field="physical_cause",
                    value="invented hardware cause",
                    evidence_id=evidence.evidence_id,
                )
            ]
            if invalid_draft
            else template_report(evidence).accepted
        )
        return ClaimDraft(claims=claims)

    monkeypatch.setattr("orbitwatch.explanation_eval.generate_draft", draft)
    output = evaluate_explanations(completed_root, "fixture", "offline-test-stub", cases=2)
    metadata = json.loads((output / "summary.json").read_text())
    summaries = {row["mode"]: row for row in metadata["summary"]}
    assert summaries["template"]["mean_latency_seconds"] is None
    assert (
        summaries["ordinary_local_llm"]["mean_latency_seconds"]
        == summaries["ordinary_llm_plus_validator"]["mean_latency_seconds"]
    )
    assert "excludes validation and rendering" in metadata["latency_scope"]
    records = json.loads((output / "records.json").read_text())
    assert len(records) == 8
    assert all(row["latency_seconds"] is None for row in records if row["mode"] == "template")
    cases = pd.read_csv(output / "cases.csv")
    assert cases.loc[cases["mode"] == "template", "latency_seconds"].isna().all()
    if invalid_draft:
        assert summaries["grounded_local_llm"]["unsupported_emitted_rate"] is None
        assert summaries["grounded_local_llm"]["mean_coverage"] == 0
    submission = create_submission(completed_root, "fixture")
    assert (submission / "Technical_Report.pdf").read_bytes().startswith(b"%PDF-")
    presentation = Presentation(submission / "Final_Presentation.pptx")
    assert len(presentation.slides) == 12
    if invalid_draft:
        assert any(
            "unsupported=n/a" in shape.text
            for slide in presentation.slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
